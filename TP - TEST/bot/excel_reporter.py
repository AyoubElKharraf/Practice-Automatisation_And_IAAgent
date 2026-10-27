"""
bot/excel_reporter.py - Advanced automated Excel reporting using OpenPyXL.
Generates an executive workbook featuring KPI cards, formatted data tables,
clickable hyperlinks, and native Excel analytical charts.
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bot.config import BotConfig
from bot.logger import BotLogger

logger = BotLogger.get_logger()


class ExcelReporter:
    """
    RPA Component that produces professional Excel workbooks
    with corporate styling, auto-adjusted column widths, and charts.
    """

    # Corporate Color Palette
    COLOR_PRIMARY_NAVY = "1E293B"      # Deep Slate / Navy
    COLOR_BLUE_ACCENT = "2563EB"       # Modern Electric Blue
    COLOR_LIGHT_BG = "F8FAFC"          # Clean table alternate
    COLOR_ALERT_BG = "DCFCE7"          # Soft Green alert highlight
    COLOR_ALERT_TEXT = "166534"        # Dark Green alert text
    COLOR_CARD_BORDER = "CBD5E1"       # Light border
    COLOR_WHITE = "FFFFFF"

    def __init__(self, config: BotConfig):
        self.config = config
        self.reports_dir = self.config.output_dir
        self.reports_dir.mkdir(parents=True, exist_ok=True)

    def generate_report(self, df: Any, kpis: Dict[str, Any]) -> Path:
        """
        Creates the complete multi-sheet Excel report and saves it.
        Returns the saved file path.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{self.config.excel_prefix}_{timestamp}.xlsx"
        filepath = self.reports_dir / filename

        logger.info(f"Generating Executive Excel Report: {filepath.name}...")
        wb = openpyxl.Workbook()

        # 1. Executive Dashboard Sheet
        ws_dash = wb.active
        ws_dash.title = "Executive Summary"
        self._build_dashboard_sheet(ws_dash, kpis, df)

        # 2. All Listings Sheet
        ws_all = wb.create_sheet(title="All Job Postings")
        self._build_listings_sheet(ws_all, df, alert_only=False)

        # 3. High-Priority Alerts Sheet
        ws_alerts = wb.create_sheet(title="Priority Alerts")
        alerts_df = df[df["is_alert"] == True] if not df.empty else df
        self._build_listings_sheet(ws_alerts, alerts_df, alert_only=True)

        # 4. Market Skills & Analytics Sheet with Chart
        ws_analytics = wb.create_sheet(title="Skill Demand Analytics")
        self._build_analytics_sheet(ws_analytics, kpis)

        # Save workbook
        wb.save(filepath)
        logger.info(f"Excel Report generated successfully -> {filepath}")
        return filepath

    def _build_dashboard_sheet(self, ws, kpis: Dict[str, Any], df: Any) -> None:
        """
        Builds the executive KPI summary page with formatted metric cards.
        """
        ws.views.sheetView[0].showGridLines = True

        # Header Title Banner
        ws.merge_cells("B2:G2")
        title_cell = ws["B2"]
        title_cell.value = f"📊 {self.config.project_name.upper()} - RPA REPORT"
        title_cell.font = Font(name="Segoe UI", size=16, bold=True, color=self.COLOR_WHITE)
        title_cell.fill = PatternFill(start_color=self.COLOR_PRIMARY_NAVY, end_color=self.COLOR_PRIMARY_NAVY, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 40

        # Subtitle
        ws.merge_cells("B3:G3")
        sub_cell = ws["B3"]
        sub_cell.value = f"Automated Bot Execution Run | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Target: {self.config.location}"
        sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 20

        # KPI Metric Cards
        cards = [
            ("TOTAL HARVESTED", str(kpis.get("total_jobs", 0)), "B5", "B6"),
            ("HIGH PRIORITY ALERTS", str(kpis.get("alert_count", 0)), "C5", "C6"),
            ("AVERAGE SALARY", f"${kpis.get('avg_salary', 0):,.0f}" if kpis.get("avg_salary") else "N/A", "D5", "D6"),
            ("TOP IN-DEMAND SKILL", str(kpis.get("top_skill", "N/A")), "E5", "E6"),
            ("TOP LOCATION", str(kpis.get("top_location", "Remote")), "F5", "F6"),
        ]

        card_border = Border(
            left=Side(style="thin", color=self.COLOR_CARD_BORDER),
            right=Side(style="thin", color=self.COLOR_CARD_BORDER),
            top=Side(style="thin", color=self.COLOR_CARD_BORDER),
            bottom=Side(style="thin", color=self.COLOR_CARD_BORDER),
        )

        for label, val, top_ref, bot_ref in cards:
            c_top = ws[top_ref]
            c_top.value = label
            c_top.font = Font(name="Segoe UI", size=8, bold=True, color="475569")
            c_top.fill = PatternFill(start_color=self.COLOR_LIGHT_BG, fill_type="solid")
            c_top.alignment = Alignment(horizontal="center", vertical="center")
            c_top.border = card_border

            c_bot = ws[bot_ref]
            c_bot.value = val
            c_bot.font = Font(name="Segoe UI", size=14, bold=True, color=self.COLOR_BLUE_ACCENT)
            c_bot.alignment = Alignment(horizontal="center", vertical="center")
            c_bot.border = card_border

        ws.row_dimensions[5].height = 22
        ws.row_dimensions[6].height = 32

        # Top 5 Recommended Roles Section
        ws["B8"].value = "⭐ Top Ranked Opportunities (Highest Match Scores)"
        ws["B8"].font = Font(name="Segoe UI", size=12, bold=True, color=self.COLOR_PRIMARY_NAVY)

        top_headers = ["Rank", "Job Title", "Company", "Location", "Est. Salary", "Match Score"]
        for col_idx, header in enumerate(top_headers, start=2):
            cell = ws.cell(row=9, column=col_idx, value=header)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_WHITE)
            cell.fill = PatternFill(start_color=self.COLOR_BLUE_ACCENT, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[9].height = 24

        if not df.empty:
            top_records = df.head(5)
            for row_idx, (_, r) in enumerate(top_records.iterrows(), start=10):
                ws.cell(row=row_idx, column=2, value=f"#{row_idx - 9}").alignment = Alignment(horizontal="center")
                ws.cell(row=row_idx, column=3, value=r.get("title", ""))
                ws.cell(row=row_idx, column=4, value=r.get("company", ""))
                ws.cell(row=row_idx, column=5, value=r.get("location", ""))
                sal = r.get("salary_estimated", 0)
                ws.cell(row=row_idx, column=6, value=f"${sal:,.0f}" if sal > 0 else "Not Disclosed").alignment = Alignment(horizontal="right")
                ws.cell(row=row_idx, column=7, value=f"{r.get('match_score', 0)}%").alignment = Alignment(horizontal="center")
                ws.row_dimensions[row_idx].height = 20

        # Auto fit column widths
        self._auto_fit_columns(ws, max_cols=8)

    def _build_listings_sheet(self, ws, df: Any, alert_only: bool = False) -> None:
        """
        Builds detailed tables with headers, zebra striping, and clickable links.
        """
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Match", "Job Title", "Company", "Location",
            "Salary (Raw)", "Est. Annual USD", "Skills Detected",
            "Source", "Method", "Action / Link"
        ]

        header_fill = PatternFill(
            start_color=self.COLOR_ALERT_TEXT if alert_only else self.COLOR_PRIMARY_NAVY,
            fill_type="solid"
        )
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_WHITE)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        if not df.empty:
            for row_idx, (_, r) in enumerate(df.iterrows(), start=2):
                bg_color = self.COLOR_ALERT_BG if r.get("is_alert") and not alert_only else (
                    self.COLOR_LIGHT_BG if row_idx % 2 == 0 else self.COLOR_WHITE
                )
                row_fill = PatternFill(start_color=bg_color, fill_type="solid")

                # Match Score
                c1 = ws.cell(row=row_idx, column=1, value=f"{r.get('match_score', 0)}%")
                c1.alignment = Alignment(horizontal="center", vertical="center")

                # Title & Company
                c2 = ws.cell(row=row_idx, column=2, value=r.get("title", ""))
                c3 = ws.cell(row=row_idx, column=3, value=r.get("company", ""))
                c4 = ws.cell(row=row_idx, column=4, value=r.get("location", ""))
                c5 = ws.cell(row=row_idx, column=5, value=r.get("salary_raw", ""))

                # Numeric Salary
                c6 = ws.cell(row=row_idx, column=6, value=r.get("salary_estimated", 0))
                c6.number_format = "$#,##0"
                c6.alignment = Alignment(horizontal="right", vertical="center")

                c7 = ws.cell(row=row_idx, column=7, value=r.get("skills_str", ""))
                c8 = ws.cell(row=row_idx, column=8, value=r.get("source", ""))
                c9 = ws.cell(row=row_idx, column=9, value=r.get("harvest_method", ""))

                # Hyperlink
                link_url = r.get("url", "")
                c10 = ws.cell(row=row_idx, column=10, value="Apply / View")
                if link_url and link_url.startswith("http"):
                    c10.hyperlink = link_url
                    c10.font = Font(name="Segoe UI", size=9, underline="single", color="2563EB")
                c10.alignment = Alignment(horizontal="center", vertical="center")

                for c_col in range(1, 11):
                    cell = ws.cell(row=row_idx, column=c_col)
                    cell.border = thin_border
                    if c_col != 10:  # preserve hyperlink font
                        cell.font = Font(name="Segoe UI", size=9)
                    if c_col != 10 or not (link_url and link_url.startswith("http")):
                        cell.fill = row_fill

                ws.row_dimensions[row_idx].height = 20

        # Enable Excel Auto-Filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df) + 1}"
        self._auto_fit_columns(ws, max_cols=len(headers))

    def _build_analytics_sheet(self, ws, kpis: Dict[str, Any]) -> None:
        """
        Builds skill frequency table and attaches a native OpenPyXL BarChart.
        """
        ws.views.sheetView[0].showGridLines = True

        ws["B2"].value = "Technical Skill Market Demand"
        ws["B2"].font = Font(name="Segoe UI", size=14, bold=True, color=self.COLOR_PRIMARY_NAVY)

        ws["B4"].value = "Skill / Keyword"
        ws["C4"].value = "Job Post Count"
        for col in ["B4", "C4"]:
            ws[col].font = Font(name="Segoe UI", size=10, bold=True, color=self.COLOR_WHITE)
            ws[col].fill = PatternFill(start_color=self.COLOR_PRIMARY_NAVY, fill_type="solid")
            ws[col].alignment = Alignment(horizontal="center", vertical="center")

        skill_dist = kpis.get("skill_distribution", {})
        current_row = 5

        for skill, count in skill_dist.items():
            c_skill = ws.cell(row=current_row, column=2, value=skill)
            c_count = ws.cell(row=current_row, column=3, value=count)
            c_skill.font = Font(name="Segoe UI", size=10)
            c_count.font = Font(name="Segoe UI", size=10)
            c_count.alignment = Alignment(horizontal="center")
            current_row += 1

        # Add Native OpenPyXL Bar Chart
        if skill_dist:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Top Required In-Demand Skills in Market"
            chart.y_axis.title = "Frequency in Harvested Roles"
            chart.x_axis.title = "Skill"
            chart.width = 16
            chart.height = 10

            data = Reference(ws, min_col=3, min_row=4, max_row=current_row - 1)
            cats = Reference(ws, min_col=2, min_row=5, max_row=current_row - 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.legend = None

            ws.add_chart(chart, "E4")

        self._auto_fit_columns(ws, max_cols=4)

    def _auto_fit_columns(self, ws, max_cols: int = 15) -> None:
        """
        Dynamically adjusts column widths based on content size.
        """
        for col in range(1, max_cols + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for cell in ws[col_letter]:
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len and len(val_str) < 50:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
